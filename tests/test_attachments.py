import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock

from desktop_buddy.attachments import ingest,model_message,read_document,snapshot
from desktop_buddy.agent import Agent
from desktop_buddy.security import Cancelled,Denied,Policy,Redactor,StopToken
from desktop_buddy.state import Store
from desktop_buddy.tools import Registry


class AttachmentTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.store=Store(self.root/'private',Redactor());self.cid=self.store.conversation()
        self.work=self.root/'work';self.work.mkdir();self.token=StopToken()
        self.registry=Registry(Policy(self.work,'execute',True,protected=[self.store.root]),self.store)
        self.registry.conversation_id=self.cid

    def tearDown(self):self.store.close();self.temp.cleanup()

    def attach(self,name='report.txt',text='Quarterly revenue: 12000'):
        source=self.root/name;source.write_text(text)
        file=ingest(self.store.root,source)
        self.store.add_message(self.cid,{'role':'user','content':'Review this report','files':[file]})
        return source,file

    def test_reads_private_snapshot_and_preserves_original(self):
        source,file=self.attach();source.write_text('Changed after attaching')
        value=self.registry.execute('attachments.read',{'id':file['id']},self.token)
        self.assertEqual(value['text'],'Quarterly revenue: 12000')
        self.assertEqual(source.read_text(),'Changed after attaching')
        self.assertNotIn('path',self.registry.execute('attachments.list',{},self.token)[0])

    def test_scope_integrity_and_cancellation(self):
        _,file=self.attach()
        with self.assertRaises(Denied):snapshot(self.store,self.store.conversation(),file['id'])
        Path(file['path']).write_text('Changed')
        with self.assertRaisesRegex(ValueError,'changed'):snapshot(self.store,self.cid,file['id'])
        self.token.stop()
        with self.assertRaises(Cancelled):self.registry.execute('attachments.read',{'id':file['id']},self.token)

    def test_copy_cannot_overwrite_or_escape_approval_or_plan(self):
        source,file=self.attach();args={'id':file['id'],'destination':'Copy.txt'}
        self.assertTrue(self.registry.execute('attachments.copy',args,self.token)['verified'])
        self.assertEqual((self.work/'Copy.txt').read_bytes(),source.read_bytes())
        with self.assertRaises(FileExistsError):self.registry.execute('attachments.copy',args,self.token)
        args['destination']=str(self.root/'Outside.txt')
        with self.assertRaises(Denied):self.registry.execute('attachments.copy',args,self.token)
        self.registry.policy.mode='plan';args['destination']='Another.txt'
        with self.assertRaises(Denied):self.registry.execute('attachments.copy',args,self.token)
        self.assertFalse((self.work/'Another.txt').exists())

    def test_references_reach_model_and_queued_files_survive_restart(self):
        _,file=self.attach();message=self.store.messages(self.cid)[0]
        wire=model_message(message)
        self.assertIn(file['id'],wire['content']);self.assertNotIn('files',wire)
        self.assertNotIn(file['path'],wire['content']);self.assertIn('files',message)
        agent=Agent(self.store,self.registry,Mock());agent.conversation_id=self.cid;agent.inbox_closed=False
        agent.say('Read this too',files=[file]);self.store.close()
        self.store=Store(self.root/'private',Redactor())
        self.assertEqual(self.store.messages(self.cid)[-1]['files'][0],file)
        self.assertEqual(snapshot(self.store,self.cid,file['id'])[0],file)

    def test_deleting_conversation_preserves_shared_copy_and_original(self):
        source,file=self.attach();second=self.store.conversation()
        self.store.add_message(second,{'role':'user','content':'Also here','files':[file]})
        self.store.delete_conversation(self.cid);self.assertTrue(Path(file['path']).exists())
        self.store.delete_conversation(second);self.assertFalse(Path(file['path']).exists())
        self.assertTrue(source.exists())

    def test_xlsx_exact_cells_formulas_and_pagination(self):
        from openpyxl import Workbook
        source=self.root/'Sales.xlsx';book=Workbook();ws=book.active;ws.title='Sales'
        ws.append(['Region','Revenue','Total']);ws.append(['West',12000,'=SUM(B2:B3)']);ws.append(['East',9000])
        book.save(source);book.close()
        file=ingest(self.store.root,source)
        self.store.add_message(self.cid,{'role':'user','content':'Review','files':[file]})
        result=self.registry.execute('attachments.read',{'id':file['id'],'sheet':'Sales','limit':2},self.token)
        self.assertEqual(result['rows'][1]['cells']['B2']['value'],12000)
        cell=result['rows'][1]['cells']['C2'];self.assertEqual(cell['formula'],'=SUM(B2:B3)')
        self.assertIsNone(cell['cached_value']);self.assertIn('not recalculated',cell['note'])
        self.assertEqual(result['next_offset'],2)
        self.assertEqual(read_document(source,offset=2)['rows'][0]['cells']['B3']['value'],9000)
        self.assertEqual(read_document(source,offset=100)['rows'],[])

    def test_word_csv_and_pdf_report_limits(self):
        from docx import Document
        from pypdf import PdfWriter
        source=self.root/'Notes.docx';doc=Document();doc.add_paragraph('First')
        table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Item';table.cell(0,1).text='12'
        doc.add_paragraph('Last');doc.save(source)
        self.assertEqual(read_document(source)['text'],'First\nItem\t12\nLast')
        source=self.root/'Sales.csv';source.write_text('Name,Amount\n"North, West",125.50\nSouth,200\n')
        result=read_document(source,offset=1,limit=1)
        self.assertEqual(result['rows'][0]['values'],['North, West','125.50']);self.assertEqual(result['next_offset'],2)
        source=self.root/'Scan.pdf';writer=PdfWriter();writer.add_blank_page(width=300,height=300)
        with source.open('wb') as output:writer.write(output)
        result=read_document(source);self.assertEqual(result['text'],'');self.assertIn('Scanned',result['note'])
        with self.assertRaisesRegex(ValueError,'page between'):read_document(source,page=2)

    def test_text_pagination_beyond_100000_and_symlink_rejection(self):
        source=self.root/'Long.txt';source.write_text('a'*130000+'end')
        self.assertEqual(read_document(source,offset=120000)['text'],'a'*10000+'end')
        link=self.root/'link.txt';link.symlink_to(source)
        with self.assertRaises(OSError):ingest(self.store.root,link)

    def test_multibyte_text_paginates_within_worker_output_limit(self):
        _,file=self.attach('International.txt','😊'*30000)
        result=self.registry.execute('attachments.read',{'id':file['id']},self.token)
        self.assertGreater(len(result['text']),1000);self.assertLess(len(result['text']),30000)
        self.assertEqual(result['next_offset'],len(result['text']))

    def test_filename_redaction_does_not_break_snapshot_location(self):
        secret='sensitive-file-title';self.store.redactor.register(secret)
        _,file=self.attach(secret+'.txt')
        stored=self.store.messages(self.cid)[0]['files'][0]
        self.assertIn('[REDACTED]',stored['name'])
        self.assertEqual(snapshot(self.store,self.cid,file['id'])[1],Path(file['path']))

